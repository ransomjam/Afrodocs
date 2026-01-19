#!/usr/bin/env python
"""
Test placeholder replacement fix
Verifies that placeholders are completely replaced, not merged with text
"""

import os
import sys
from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'pattern-formatter/backend'))
from coverpage_generator import generate_cover_page

def test_placeholder_replacement():
    """Test that placeholders are completely replaced"""
    print("\n" + "="*70)
    print("TEST: Placeholder Replacement - No Merging")
    print("="*70)
    
    test_data = {
        'documentType': 'Internship Report',
        'institution': 'Test Institution',
        'faculty': 'Faculty of Science',
        'department': 'Computer Science',
        'studentName': 'Test Student',
        'studentId': 'TEST001',
        'title': 'Test Practicum Report',
        'supervisor': 'Dr. Test Supervisor',
        'coSupervisor': 'Prof. Co Supervisor',
        'date': '2026-01-15'
    }
    
    try:
        output_path, error = generate_cover_page(test_data)
        
        if error:
            print(f"[FAIL] Error: {error}")
            return False
        
        if not os.path.exists(output_path):
            print(f"[FAIL] Cover page not created")
            return False
        
        # Read the document and check for merged placeholders
        doc = Document(output_path)
        
        # Check paragraphs
        all_text = []
        for p in doc.paragraphs:
            all_text.append(p.text)
        
        # Check textboxes
        if doc.element.body is not None:
            for txbx in doc.element.body.iter(qn('w:txbxContent')):
                for p in txbx.iter(qn('w:p')):
                    full_p_text = ''
                    for r in p.iter(qn('w:r')):
                        for t in r.iter(qn('w:t')):
                            if t.text:
                                full_p_text += t.text
                    all_text.append(full_p_text)
        
        full_document = '\n'.join(all_text)
        
        # Check for merged placeholder issues
        problems = []
        
        # Look for patterns like {{XXXvalue}} or {{XXXVvalue}} which indicate merging
        if '{{' in full_document or '}}' in full_document:
            # Find any remaining placeholder markers
            import re
            remaining = re.findall(r'\{\{[^}]*\}\}', full_document)
            if remaining:
                problems.append(f"Found remaining placeholders: {remaining}")
        
        # Look for merged patterns like "{{REPORT TITLEPRACTICUM REPORT}}"
        if 'REPORT TITLE' in full_document and 'PRACTICUM' in full_document:
            # Check if they're merged
            if 'REPORT TITLEPRACTICUM' in full_document:
                problems.append("Found merged pattern: REPORT TITLEPRACTICUM")
        
        if 'Schoo/Faculty' in full_document and 'Faculty of Science' in full_document:
            if 'Schoo/FacultyFaculty of Science' in full_document:
                problems.append("Found merged pattern: Schoo/FacultyFaculty of Science")
        
        if problems:
            print("[FAIL] Placeholder merging detected:")
            for issue in problems:
                print(f"  - {issue}")
            print(f"\nGenerated: {os.path.basename(output_path)}")
            return False
        else:
            print("[PASS] No merged placeholder patterns detected")
            print("[PASS] All placeholders properly replaced")
            print(f"Generated: {os.path.basename(output_path)}")
            
            # Show some actual content
            print("\nDocument content preview:")
            for i, line in enumerate(all_text[:10]):
                if line.strip():
                    print(f"  {line[:80]}")
            
            return True
        
    except Exception as e:
        print(f"[FAIL] Exception: {str(e)}")
        import traceback
        traceback.print_exc()
        return False


if __name__ == '__main__':
    result = test_placeholder_replacement()
    sys.exit(0 if result else 1)
