"""Debug why placeholders aren't being replaced"""
import os
import sys
from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'pattern-formatter', 'backend'))
from coverpage_generator import generate_cover_page

# Simple test data
test_data = {
    'university': 'Bamenda',
    'documentType': 'Dissertation',
    'studentName': 'Jam Ransom',
    'studentId': 'UBa25PP227',
    'institution': 'College of Technology',
    'faculty': 'Faculty of Engineering and Technology',
    'department': 'Electrical and Electronic Engineering',
    'level': '400 Level',
    'supervisor': 'Dr. Sample',
    'academicSupervisor': 'Dr. Sample',
    'fieldSupervisor': '',
    'title': 'SAMPLE THESIS',
    'monthYear': 'January 2026'
}

print("Testing placeholder replacement...")
result = generate_cover_page(test_data)
doc_path = result[0] if isinstance(result, tuple) else result

if doc_path and os.path.exists(doc_path):
    print(f"Document generated: {os.path.basename(doc_path)}")
    doc = Document(doc_path)
    
    print("\n=== CHECKING FOR UNREPLACED PLACEHOLDERS ===")
    
    # Check paragraphs
    print("\nParagraphs:")
    placeholder_count = 0
    for i, para in enumerate(doc.paragraphs):
        if '{{' in para.text:
            print(f"  Para {i}: {para.text[:80]}")
            placeholder_count += len([p for p in para.text.split() if '{{' in p])
    
    # Check textboxes
    print("\nTextboxes:")
    if doc.element.body is not None:
        for txbx in doc.element.body.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}txbxContent'):
            for p in txbx.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                para_text = ""
                for r in p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r'):
                    for t in r.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
                        if t.text:
                            para_text += t.text
                if '{{' in para_text:
                    print(f"  Textbox: {para_text[:80]}")
                    placeholder_count += len([p for p in para_text.split() if '{{' in p])
    
    print(f"\nTotal unreplaced placeholders: {placeholder_count}")
    
    if placeholder_count > 0:
        print("\n[ISSUE] Placeholders not being replaced!")
    else:
        print("\n[OK] All placeholders replaced")
else:
    print(f"ERROR: {result[1]}")
